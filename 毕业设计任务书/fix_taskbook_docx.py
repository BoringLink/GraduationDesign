#!/usr/bin/env python3
"""根据《任务书-参考》版式修订《22010210林柏伟任务书》docx 表格内容。"""
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "_conv" / "22010210林柏伟任务书.docx"
OUT = ROOT / "_conv" / "22010210林柏伟任务书_优化.docx"

SECTION2 = """1.任务说明：
本项目旨在设计并实现一套面向中小学编程学习场景的智能编程学习平台，围绕课程学习、在线编程、学习过程记录与大模型驱动的智能辅学等核心教学环节开展系统建设。
（1）业务功能：完成用户认证与权限控制，支持课程与知识点管理、学习任务与在线编程实践、师生消息互动、教学资源管理以及学习行为记录与事件追踪等功能；
（2）前端实现：采用 Nuxt 3 与 Vue 3 构建 Web 前端，结合 TypeScript、Tailwind CSS 与 DaisyUI 完成页面与组件化界面开发；
（3）后端与数据层：采用 FastAPI 与 SQLAlchemy 提供 REST 接口与业务逻辑，使用 SQL Server 持久化业务数据，引入 Redis 承担缓存与热点加速，教学资源使用对象存储进行统一管理与分发；
（4）在线编程与安全：学生侧 Python 代码在前端通过 Pyodide 于浏览器沙箱中执行，后端聚焦鉴权、审计日志与课程/任务等业务服务，降低服务端任意代码执行风险；
（5）大模型辅学：集成大语言模型服务接口，支持编程问答、代码与报错解释、任务分步引导以及面向中小学生认知特点的输出约束与引导策略；
（6）系统验证：完成需求分析、总体设计、核心模块实现与功能/非功能测试，撰写毕业设计说明书并完成学位论文规定的图表与技术文档。

2.任务要求：
（1）能够积极、主动且认真地学习并完成任务；
（2）具备 Web 应用开发基础，熟悉 Python、JavaScript/TypeScript 及常用前后端框架；
（3）关注中小学编程教育场景下的可用性、安全性与数据合规，按要求完成开题、中期与论文撰写环节。

3.参考文献：
Wang S, Xu T, Li H, et al. Large Language Models for Education: A Survey and Outlook[EB/OL]. arXiv:2403.18105, 2024.
Raihan N, Siddiq M L, Santos J C S, et al. Large Language Models in Computer Science Education: A Systematic Literature Review[C]// Proceedings of the 56th ACM Technical Symposium on Computer Science Education. New York: ACM, 2025: 938-944.
Pitts G, Hridi A P, Lekshmi-Narayanan A B. A Survey of LLM-Based Applications in Programming Education: Balancing Automation and Human Oversight[EB/OL]. arXiv:2510.03719, 2025.
Tang B, Liang J, Hu W, et al. Enhancing Programming Performance, Learning Interest, and Self-Efficacy: The Role of Large Language Models in Middle School Education[J]. Systems, 2025, 13(7): 555.
Yan Y M, Chen C Q, Hu Y B, et al. LLM-based collaborative programming: impact on students' computational thinking and self-efficacy[J]. Humanities and Social Sciences Communications, 2025, 12(1): 149.
Pirzado F A, Ahmed A, Mendoza-Urdiales R A, et al. Navigating the Pitfalls: Analyzing the Behavior of LLMs as a Coding Assistant for Computer Science Students[J]. IEEE Access, 2024, 12: 112605-112625.
Chen S J, Shan X, Liu Z M, et al. Employing large language models to enhance K-12 students' programming debugging skills, computational thinking, and self-efficacy[J]. Educational Technology & Society, 2025, 28(2): 259-278.
Yim I H Y, Su J. Artificial intelligence (AI) learning tools in K-12 education: A scoping review[J]. Journal of Computers in Education, 2024, DOI: 10.1007/s40692-023-00304-9.
Tsai M J, Liang J C, Hsu C Y. The Computational Thinking Scale for Computer Literacy Education[J]. Journal of Educational Computing Research, 2021, 59(4): 579-602.
Sun D, Zhu C, Xu F, et al. Transitioning from introductory to professional programming in secondary education: Comparing learners' computational thinking skills, behaviors, and attitudes[J]. Journal of Educational Computing Research, 2023, DOI: 10.1177/07356331231204653.
Touretzky D, Gardner-McCune C, Martin F, et al. Envisioning AI for K-12: What Should Every Child Know about AI?[C]// Proceedings of the AAAI Conference on Artificial Intelligence. 2019, 33(1): 9795-9799.
姚佳佳，李艳，潘金晶，等。同伴对话反馈对大学生在线深度学习的影响研究[J]. 华东师范大学学报（教育科学版）, 2022, 40(3): 112-126.
教育部基础教育教学指导委员会。中小学生生成式人工智能使用指南（2025 年版）[Z]. 北京：教育部，2025.
"""


def main() -> None:
    doc = Document(SRC)
    assert len(doc.tables) >= 8, "表格数量异常"

    doc.tables[2].rows[0].cells[0].text = SECTION2.strip() + "\n"

    # 进度表第 4 行时间列笔误 202512.25
    sched = doc.tables[4]
    sched.rows[5].cells[1].text = "2025.12.25-2025.12.29"

    sig = doc.tables[5].rows[0].cells
    sig[3].text = "2025"
    sig[5].text = "11"
    sig[7].text = "10"

    rev = doc.tables[7].rows[0].cells
    rev[0].text = "研究所（专业）负责人       "
    rev[1].text = "     "
    rev[3].text = "2025"
    rev[5].text = "11"
    rev[7].text = "12"

    doc.save(OUT)
    print("written", OUT)


if __name__ == "__main__":
    main()
