#!/usr/bin/env python3
"""fix_thesis_formatting_transplant.py

将“论文重写版_优化后.docx”的内容移植到“理工类论文模板（自动生成目录版）.docx”的结构中，
生成符合模板格式的输出文档。

约束（来自任务说明）：
- 不修改 thesis 源文档
- 不修改 CodingPlatformBak/
- 不改模板 Normal 样式定义
- 删除段落范围时需同时删除范围内的表格(w:tbl)
- 插入到指定位置必须使用 lxml 的 addnext()（不依赖 add_paragraph 的末尾追加）
"""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


TEMPLATE_PATH = Path(
    "/Users/tk/Documents/杭州电子科技大学/毕业设计/毕业设计相关材料/07-理工类论文模板（自动生成目录版）.docx"
)
THESIS_PATH = Path(
    "/Users/tk/Documents/杭州电子科技大学/毕业设计/论文重写版_优化后.docx"
)
OUTPUT_PATH = Path(
    "/Users/tk/Documents/杭州电子科技大学/毕业设计/论文重写版_格式修正.docx"
)


THESIS_TITLE = "基于大模型的中小学生智能编程学习平台设计与实现"


ORAL_PREFIXES = [
    "回顾以上讨论，在多数情况下，",
    "在上述结论的延长线上，",
    "这里有一个值得在意的细节，",
    "这里有一个值得着眼的细节，",
    "就这一点延伸而言，",
    "回顾以上讨论，",
    "从实现角度来看，",
    "从办到角度来看，",
    "就办到角度而言，",
    "在系统设计层面，",
    "就技术实现而言，",
    "就技术办到而言，",
    "就技术达成而言，",
    "实现角度的层面上，",
    "从实际应用角度，",
    "基于上述分析，",
    "糅合具体场景，",
    "在系统设计层面，",
    "就技术实现而言，",
    "进一步地，",
    "展开来说，",
    "这也提示我们，",
    "最初，",
    "第二，",
    "第二。",
    "另外还有，",
    "同时，",
    "第一，",
    "除上述之外，",
    "同样，",
    "深入地，",
    "以此为起点，",
    "逐一来看，",
]


def _norm_ws(s: str) -> str:
    return " ".join((s or "").split())


def _p_text(p_elm) -> str:
    parts: list[str] = []
    for n in p_elm.iter():
        if n.tag == f"{{{W_NS}}}t" and n.text:
            parts.append(n.text)
    return "".join(parts)


def _has_sectPr_in_p(p_elm) -> bool:
    for n in p_elm.iter():
        if n.tag == f"{{{W_NS}}}sectPr":
            return True
    return False


def _has_page_break(p_elm) -> bool:
    for n in p_elm.iter():
        if n.tag == f"{{{W_NS}}}br" and n.get(f"{{{W_NS}}}type") == "page":
            return True
    return False


def _body_children(body) -> list:
    return list(body.iterchildren())


def _find_child_idx(body, pred) -> int:
    for idx, el in enumerate(_body_children(body)):
        if pred(idx, el):
            return idx
    raise RuntimeError("未找到目标元素")


def _find_paragraph_child_idx_by_text(body, wanted_norm_text: str) -> int:
    wanted_norm_text = _norm_ws(wanted_norm_text)

    def pred(_idx, el):
        if el.tag != f"{{{W_NS}}}p":
            return False
        return _norm_ws(_p_text(el)) == wanted_norm_text

    return _find_child_idx(body, pred)


def _remove_children_inclusive(body, start_child_idx: int, end_child_idx: int) -> None:
    children = _body_children(body)
    if start_child_idx < 0 or end_child_idx >= len(children) or start_child_idx > end_child_idx:
        raise ValueError("remove range out of bounds")

    for i in range(end_child_idx, start_child_idx - 1, -1):
        el = children[i]
        # 只允许删除段落或表格，避免误删 body 末尾 sectPr
        if el.tag in (f"{{{W_NS}}}p", f"{{{W_NS}}}tbl"):
            body.remove(el)


def _insert_after(anchor_elm, new_elms: list) -> None:
    cur = anchor_elm
    for el in new_elms:
        cur.addnext(el)
        cur = el


def _clean_text(s: str) -> str:
    s = (s or "").strip()

    # 先做全局纠错，使“办到”等变体能命中统一的口语前缀列表
    s = s.replace("办到", "实现")
    s = s.replace("呈上", "提供")
    s = s.replace("扶持", "支持")

    # 去口语前缀（可多次）
    prefixes = sorted(set(ORAL_PREFIXES), key=len, reverse=True)
    while True:
        removed = False
        for pref in prefixes:
            if s.startswith(pref):
                s = s[len(pref) :].lstrip()
                removed = True
                break
        if not removed:
            break

    # 再纠错一次（防止前缀去除后新出现可替换片段）
    s = s.replace("办到", "实现")
    s = s.replace("呈上", "提供")
    s = s.replace("扶持", "支持")
    return s


def _strip_keywords_suffix_en(s: str) -> str:
    # 去除尾部追加的 "Key words：..."
    m = re.search(r"\bKey\s+words\s*[:：]", s)
    if not m:
        return s
    return s[: m.start()].rstrip()


def _make_p(style_id: str | None = None, line_400_exact: bool = False):
    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    if style_id:
        pStyle = OxmlElement("w:pStyle")
        pStyle.set(qn("w:val"), style_id)
        pPr.append(pStyle)
    if line_400_exact:
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:line"), "400")
        spacing.set(qn("w:lineRule"), "exact")
        pPr.append(spacing)
    p.append(pPr)
    return p


def _add_text_run(p, text: str):
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    p.append(r)


def _add_formatted_run(
    p,
    text: str,
    *,
    bold: bool = False,
    sz: int = 24,
    font_east_asia: str | None = None,
    font_ascii: str | None = None,
):
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    if bold:
        b = OxmlElement("w:b")
        rPr.append(b)
    if font_east_asia or font_ascii:
        rFonts = OxmlElement("w:rFonts")
        if font_ascii:
            rFonts.set(qn("w:ascii"), font_ascii)
            rFonts.set(qn("w:hAnsi"), font_ascii)
        if font_east_asia:
            rFonts.set(qn("w:eastAsia"), font_east_asia)
        rPr.append(rFonts)
    sz_elm = OxmlElement("w:sz")
    sz_elm.set(qn("w:val"), str(sz))
    rPr.append(sz_elm)
    szcs_elm = OxmlElement("w:szCs")
    szcs_elm.set(qn("w:val"), str(sz))
    rPr.append(szcs_elm)
    r.append(rPr)

    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    p.append(r)


def _make_page_break_p() -> object:
    p = _make_p(style_id="Normal", line_400_exact=True)
    r = OxmlElement("w:r")
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    r.append(br)
    p.append(r)
    return p


def _set_cover_table(doc: Document) -> None:
    t = doc.tables[0]
    mapping = {
        "题    目": THESIS_TITLE,
        "专    业": "计算机科学与技术",
        "班    级": "22052323",
        "学    号": "22010210",
        "学生姓名": "林柏伟",
        "指导教师": "匡振中",
    }
    for row in t.rows:
        key = row.cells[0].text.replace("\n", " ").strip()
        if key not in mapping:
            continue
        cell = row.cells[1]
        value = mapping[key]
        # 尽量保留单元格原有段落/格式：只改第一个 run 的文本
        if not cell.paragraphs:
            cell.text = value
            continue
        first_p = cell.paragraphs[0]
        if first_p.runs:
            first_p.runs[0].text = value
            for r in first_p.runs[1:]:
                r.text = ""
        else:
            first_p.add_run(value)
        for p in cell.paragraphs[1:]:
            p.text = ""


def _replace_honesty_title(doc: Document) -> None:
    body = doc.element.body
    for el in _body_children(body):
        if el.tag != f"{{{W_NS}}}p":
            continue
        txt = _p_text(el)
        if "我谨在此承诺：本人所写的毕业论文《" not in txt:
            continue
        # 只替换《 》中的占位空格
        new_txt = re.sub(r"《\s*》", f"《{THESIS_TITLE}》", txt)
        if new_txt == txt:
            # 兼容模板里用大量空格占位
            new_txt = txt.replace("《                    》", f"《{THESIS_TITLE}》")
        # 该段落模板为单个 w:t，直接改文本即可
        for n in el.iter():
            if n.tag == f"{{{W_NS}}}t":
                n.text = new_txt
                return
    raise RuntimeError("未找到诚信承诺段落")


def _rebuild_abstract_cn(doc: Document, thesis: Document) -> None:
    body = doc.element.body
    idx_abstract = _find_paragraph_child_idx_by_text(body, "摘 要")
    idx_abs_en = _find_paragraph_child_idx_by_text(body, "ABSTRACT")

    # 删除 摘要标题 与 ABSTRACT 标题之间的所有内容（包含模板关键词/空行/分页符）
    if idx_abstract + 1 <= idx_abs_en - 1:
        _remove_children_inclusive(body, idx_abstract + 1, idx_abs_en - 1)

    # 插入 thesis P2-P5
    new_elms: list = []
    for i in range(2, 6):
        s = _clean_text(thesis.paragraphs[i].text)
        p = _make_p(style_id="Normal", line_400_exact=True)
        _add_formatted_run(p, s, sz=24)
        new_elms.append(p)

    # 关键词段
    kw_p = _make_p(style_id="Normal", line_400_exact=True)
    _add_formatted_run(kw_p, "关键词：", bold=True, sz=24, font_east_asia="黑体")
    _add_formatted_run(
        kw_p,
        "大语言模型；编程教育；智能辅学；在线编程；教学平台",
        sz=24,
    )
    new_elms.append(kw_p)

    # 9个空行 + 分页符段落
    for _ in range(9):
        new_elms.append(_make_p(style_id="Normal", line_400_exact=True))
    new_elms.append(_make_page_break_p())

    # 插入到 P19 (摘要标题) 之后
    anchor = _body_children(body)[idx_abstract]
    _insert_after(anchor, new_elms)


def _rebuild_abstract_en(doc: Document, thesis: Document) -> None:
    body = doc.element.body
    idx_abs_en = _find_paragraph_child_idx_by_text(body, "ABSTRACT")

    # 找到 ABSTRACT 后的 sectPr 段落（模板 P42），必须保留
    children = _body_children(body)
    idx_sect_p = None
    for i in range(idx_abs_en + 1, len(children)):
        el = children[i]
        if el.tag == f"{{{W_NS}}}p" and _has_sectPr_in_p(el):
            idx_sect_p = i
            break
    if idx_sect_p is None:
        raise RuntimeError("未找到 ABSTRACT 后的分节符段落")

    # 删除 ABSTRACT 标题与 sectPr 段落之间的所有内容
    if idx_abs_en + 1 <= idx_sect_p - 1:
        _remove_children_inclusive(body, idx_abs_en + 1, idx_sect_p - 1)

    new_elms: list = []
    for i in range(7, 11):
        s = _clean_text(thesis.paragraphs[i].text)
        if i == 10:
            s = _strip_keywords_suffix_en(s)
        p = _make_p(style_id="Normal", line_400_exact=True)
        _add_formatted_run(p, s, sz=24, font_ascii="Times New Roman")
        new_elms.append(p)

    kw = "large language model; programming education; intelligent tutoring; online programming; learning platform"
    kw_p = _make_p(style_id="Normal", line_400_exact=True)
    _add_formatted_run(kw_p, "Key words：", bold=True, sz=24, font_ascii="Times New Roman")
    _add_formatted_run(kw_p, kw, sz=24, font_ascii="Times New Roman")
    new_elms.append(kw_p)

    anchor = _body_children(body)[idx_abs_en]
    _insert_after(anchor, new_elms)


def _build_body_from_thesis(thesis: Document) -> list:
    out: list = []

    def add_body_para(text: str):
        p = _make_p(style_id="2")
        _add_text_run(p, text)
        out.append(p)

    def add_heading(style_id: str, text: str, add_page_break_before: bool):
        if add_page_break_before:
            out.append(_make_page_break_p())
        p = _make_p(style_id=style_id)
        _add_text_run(p, text)
        out.append(p)

    # thesis 正文 1-7 章
    first_h1_inserted = False
    for i in range(11, 219):
        p_src = thesis.paragraphs[i]
        sid = p_src.style.style_id
        text = _clean_text(p_src.text)
        if sid == "Heading1":
            # 模板要求：除第一章开头外，Heading1 前要加分页段
            add_pb = first_h1_inserted
            if "软件做到" in text:
                text = text.replace("软件做到", "软件实现")
            add_heading("Heading1", text, add_pb)
            first_h1_inserted = True
        elif sid == "Heading2":
            add_heading("Heading2", text, False)
        elif sid == "Heading3":
            add_heading("Heading3", text, False)
        else:
            add_body_para(text)

    # 8 结论（thesis P219-P222）
    add_heading("Heading1", _clean_text(thesis.paragraphs[219].text), True)
    for i in (220, 221, 222):
        t = _clean_text(thesis.paragraphs[i].text)
        if i == 220:
            # 额外要求：去前缀 + “做到了”->“实现了”
            t = t.replace("做到了", "实现了")
        if i == 222:
            t = t.replace("做到", "实现")
        add_body_para(t)

    # 参考文献（thesis P223-P245）
    add_heading("Heading1", _clean_text(thesis.paragraphs[223].text), True)
    refs: list[str] = []
    for i in range(224, 246):
        t = _clean_text(thesis.paragraphs[i].text)
        if not t:
            continue
        if i == 238:
            # 特殊：段首混入“16”
            t = re.sub(r"^\s*\d+(?=[\u4e00-\u9fff])", "", t)
        if i == 243:
            # 特殊：三个参考文献合并在一个段落，以“。”分隔
            parts = [seg.strip() for seg in t.split("。") if seg.strip()]
            for seg in parts:
                seg = seg.rstrip("。").strip()
                refs.append(seg + "。")
            continue
        refs.append(t)

    for idx, ref in enumerate(refs, start=1):
        add_body_para(f"[{idx}] {ref}")

    # 致谢（thesis P246-P247）
    add_heading("Heading1", _clean_text(thesis.paragraphs[246].text), True)
    add_body_para(_clean_text(thesis.paragraphs[247].text))

    return out


def _replace_body_main(doc: Document, thesis: Document) -> None:
    body = doc.element.body
    children = _body_children(body)

    # 定位目录页："目 录" -> 下一个包含分页符的空段(P44)
    idx_toc = _find_paragraph_child_idx_by_text(body, "目 录")
    children = _body_children(body)
    idx_toc_pb = None
    for i in range(idx_toc + 1, len(children)):
        el = children[i]
        if el.tag == f"{{{W_NS}}}p" and _has_page_break(el):
            idx_toc_pb = i
            break
    if idx_toc_pb is None:
        raise RuntimeError("未找到目录后的分页段落")

    # 删除模板正文：从第一个 "1 引言" 开始，到 body 末尾 sectPr 之前
    children = _body_children(body)
    last_idx = len(children) - 1
    if children[last_idx].tag != f"{{{W_NS}}}sectPr":
        raise RuntimeError("模板 body 末尾不是 sectPr，无法安全处理")

    def is_intro_p(_idx, el):
        if el.tag != f"{{{W_NS}}}p":
            return False
        return _norm_ws(_p_text(el)) == "1 引言"

    idx_intro = None
    for i in range(idx_toc_pb + 1, last_idx):
        if is_intro_p(i, children[i]):
            idx_intro = i
            break
    if idx_intro is None:
        raise RuntimeError("未找到模板正文起始段落 '1 引言'")

    if idx_intro <= last_idx - 1:
        _remove_children_inclusive(body, idx_intro, last_idx - 1)

    # 插入 thesis 正文（以及结论/参考文献/致谢）
    new_body_elms = _build_body_from_thesis(thesis)
    anchor = _body_children(body)[idx_toc_pb]
    _insert_after(anchor, new_body_elms)


def main() -> None:
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(str(TEMPLATE_PATH))
    if not THESIS_PATH.exists():
        raise FileNotFoundError(str(THESIS_PATH))

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)

    # 1) 复制模板到输出路径
    shutil.copy2(TEMPLATE_PATH, OUTPUT_PATH)

    # 2) 打开文档
    out_doc = Document(str(OUTPUT_PATH))
    thesis_doc = Document(str(THESIS_PATH))

    # 3) 填充封面表
    _set_cover_table(out_doc)

    # 4) 替换诚信承诺标题占位
    _replace_honesty_title(out_doc)

    # 5) 重建中文摘要
    _rebuild_abstract_cn(out_doc, thesis_doc)

    # 6) 重建英文摘要
    _rebuild_abstract_en(out_doc, thesis_doc)

    # 7) 正文整体替换（保留目录页、分节符与 body 最终 sectPr）
    _replace_body_main(out_doc, thesis_doc)

    # 保存
    out_doc.save(str(OUTPUT_PATH))
    print(f"OK: 已生成 {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
