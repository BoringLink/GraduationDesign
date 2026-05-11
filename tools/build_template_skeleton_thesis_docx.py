from __future__ import annotations

import copy
import re
import shutil
import subprocess
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

import build_full_thesis_docx as base


ROOT = Path(__file__).resolve().parents[1]
THESIS = ROOT / "latex" / "thesis"
TEMPLATE = ROOT / "latex" / "hdu-template.docx"
WORK = Path("/private/tmp/hdu-template-skeleton-word")
OUT = ROOT / "论文模板骨架迁移版.docx"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def delete_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def replace_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def insert_text_paragraph_before(paragraph, text: str, style_name: str = "Body Text"):
    p = paragraph.insert_paragraph_before(text)
    try:
        p.style = paragraph.part.document.styles[style_name]
    except Exception:
        pass
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = Pt(20)
    for run in p.runs:
        base.set_run_font(run, "宋体", "Times New Roman", 12)
    return p


def fill_template_front(path: Path) -> None:
    doc = Document(path)
    meta = base.metadata()
    cover_values = [
        ("题    目", meta["ThesisTitleCn"]),
        ("站点名称", meta["ThesisSite"]),
        ("专    业", meta["ThesisMajor"]),
        ("班    级", meta["ThesisClass"]),
        ("学    号", meta["ThesisStudentId"]),
        ("学生姓名", meta["ThesisAuthor"]),
        ("指导教师", meta["ThesisAdvisor"]),
        ("完成日期", meta["ThesisDate"]),
    ]
    table = doc.tables[0]
    for row, (label, value) in zip(table.rows, cover_values):
        replace_paragraph_text(row.cells[0].paragraphs[0], label)
        replace_paragraph_text(row.cells[1].paragraphs[0], value)
        for cell in row.cells:
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.first_line_indent = Pt(0)
                for run in para.runs:
                    base.set_run_font(run, "楷体", "Times New Roman", 15)

    for para in doc.paragraphs:
        text = para.text.strip()
        if text == "（2022届）":
            replace_paragraph_text(para, "（2026届）")
        elif "本人所写的毕业论文《XXXXXXXX》" in text:
            replace_paragraph_text(
                para,
                f"我谨在此承诺：本人所写的毕业论文《{meta['ThesisTitleCn']}》均系本人独立完成，没有抄袭行为，凡涉及其他作者的观点和材料，均作了注释，若有不实，后果由本人承担。",
            )

    cn_body, cn_kw, en_body, en_kw = base.abstract_parts()
    paras = doc.paragraphs

    def find_text(target: str) -> int:
        for idx, para in enumerate(doc.paragraphs):
            if para.text.strip() == target:
                return idx
        raise ValueError(target)

    cn_title = find_text("摘    要")
    cn_kw_idx = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip().startswith("关键词"))
    for idx in range(cn_kw_idx - 1, cn_title, -1):
        delete_paragraph(doc.paragraphs[idx])
    cn_kw_para = next(p for p in doc.paragraphs if p.text.strip().startswith("关键词"))
    for text in cn_body:
        insert_text_paragraph_before(cn_kw_para, text)
    replace_paragraph_text(cn_kw_para, f"关键词：{cn_kw}")
    cn_kw_para.paragraph_format.first_line_indent = Pt(0)
    for run in cn_kw_para.runs:
        base.set_run_font(run, "宋体", "Times New Roman", 12)

    en_title_idx = find_text("ABSTRACT")
    # Remove filler between Chinese keywords and English abstract, then preserve separation with a page break.
    cn_kw_idx = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip().startswith("关键词"))
    for idx in range(en_title_idx - 1, cn_kw_idx, -1):
        if not doc.paragraphs[idx].text.strip():
            delete_paragraph(doc.paragraphs[idx])
    en_title = next(p for p in doc.paragraphs if p.text.strip() == "ABSTRACT")
    br = en_title.insert_paragraph_before("")
    br.paragraph_format.first_line_indent = Pt(0)
    br.add_run().add_break(WD_BREAK.PAGE)

    en_kw_para = next(p for p in doc.paragraphs if p.text.strip().startswith("Key words"))
    en_title_idx = find_text("ABSTRACT")
    en_kw_idx = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip().startswith("Key words"))
    for idx in range(en_kw_idx - 1, en_title_idx, -1):
        delete_paragraph(doc.paragraphs[idx])
    for text in en_body:
        p = insert_text_paragraph_before(en_kw_para, text)
        for run in p.runs:
            base.set_run_font(run, "Times New Roman", "Times New Roman", 12)
    replace_paragraph_text(en_kw_para, f"Key words：{en_kw}")
    en_kw_para.paragraph_format.first_line_indent = Pt(0)
    for run in en_kw_para.runs:
        base.set_run_font(run, "Times New Roman", "Times New Roman", 12)

    toc_title = next(p for p in doc.paragraphs if p.text.strip() == "目    录")
    en_kw_idx = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip().startswith("Key words"))
    toc_idx = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "目    录")
    for idx in range(toc_idx - 1, en_kw_idx, -1):
        if not doc.paragraphs[idx].text.strip():
            delete_paragraph(doc.paragraphs[idx])
    br = toc_title.insert_paragraph_before("")
    br.paragraph_format.first_line_indent = Pt(0)
    br.add_run().add_break(WD_BREAK.PAGE)

    # Replace stale TOC result lines, but keep the template TOC heading location.
    first_body_idx = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip().endswith("引言") and p.style.name == "Heading 1")
    toc_idx = next(i for i, p in enumerate(doc.paragraphs) if p.text.strip() == "目    录")
    for idx in range(first_body_idx - 1, toc_idx, -1):
        delete_paragraph(doc.paragraphs[idx])
    first_body = next(p for p in doc.paragraphs if p.text.strip().endswith("引言") and p.style.name == "Heading 1")
    toc_field = first_body.insert_paragraph_before("")
    toc_field.paragraph_format.first_line_indent = Pt(0)
    base.add_toc(toc_field)
    br = first_body.insert_paragraph_before("")
    br.paragraph_format.first_line_indent = Pt(0)
    br.add_run().add_break(WD_BREAK.PAGE)

    doc.save(path)
    base.patch_update_fields(path)


def build_content_source() -> Path:
    order, refs = base.parse_bbl()
    draft = base.replace_citations(base.read(THESIS / "chapters" / "draft.tex"), order)
    draft = base.prefer_word_safe_images(draft)
    draft = re.sub(r"\\label\{[^}]+\}", "", draft)
    refs_tex = ["\\chapter*{参考文献}"] + [f"\\noindent [{idx}] {ref}\n" for idx, ref in enumerate(refs, 1)]
    source = "\n\n".join([draft, "\n".join(refs_tex), base.backmatter_after_refs(order)])
    WORK.mkdir(parents=True, exist_ok=True)
    out = WORK / "body-backmatter.tex"
    out.write_text(source, encoding="utf-8")
    return out


def format_content_docx(path: Path) -> None:
    doc = Document(path)
    chapter = section = subsection = 0
    in_refs = False
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if para.style.name == "Heading 1":
            if text in {"参考文献", "致谢", "附录"}:
                para.text = text
                in_refs = text == "参考文献"
            else:
                chapter += 1
                section = subsection = 0
                para.text = f"{chapter}  {text}"
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.first_line_indent = Pt(0)
            if chapter > 1 or text in {"参考文献", "致谢", "附录"}:
                br = para.insert_paragraph_before("")
                br.paragraph_format.first_line_indent = Pt(0)
                br.add_run().add_break(WD_BREAK.PAGE)
        elif para.style.name == "Heading 2":
            if not text.startswith("附录 "):
                section += 1
                subsection = 0
                para.text = f"{chapter}.{section}  {text}"
            para.paragraph_format.first_line_indent = Pt(0)
        elif para.style.name == "Heading 3":
            subsection += 1
            para.text = f"{chapter}.{section}.{subsection}  {text}"
            para.paragraph_format.first_line_indent = Pt(0)
        if re.match(r"^[图表]\s*\d", text):
            para.style = doc.styles["Caption"]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.first_line_indent = Pt(0)
        for run in para.runs:
            if para.style.name.startswith("Heading"):
                size = 16 if para.style.name == "Heading 1" else 15 if para.style.name == "Heading 2" else 14
                base.set_run_font(run, "黑体", "Times New Roman", size, True)
            elif para.style.name == "Caption":
                base.set_run_font(run, "宋体", "Times New Roman", 10.5, True)
            elif in_refs:
                base.set_run_font(run, "宋体", "Times New Roman", 10.5)
            else:
                base.set_run_font(run, "宋体", "Times New Roman", 12)

    content_width = base.section_content_width_dxa(doc.sections[0])
    for table in doc.tables:
        cols = len(table.columns)
        if cols == 2:
            widths = base.column_widths_from_weights([1.2, 3.8], content_width)
        elif cols == 5:
            widths = base.column_widths_from_weights([0.9, 1.5, 2.4, 2.4, 0.8], content_width)
        else:
            widths = base.column_widths_from_weights([1] * cols, content_width)
        base.apply_table_geometry(table, widths, table_width_dxa=sum(widths), indent_dxa=0)
    doc.save(path)


def text_of(elem) -> str:
    return "".join(elem.xpath(".//w:t/text()", namespaces={"w": W_NS}))


def merge_body_into_template(front_docx: Path, content_docx: Path, out_docx: Path) -> None:
    temp = WORK / "merge"
    if temp.exists():
        shutil.rmtree(temp)
    temp.mkdir(parents=True)
    with zipfile.ZipFile(front_docx) as z:
        z.extractall(temp)

    parser = etree.XMLParser(remove_blank_text=False)
    doc_xml_path = temp / "word" / "document.xml"
    base_tree = etree.parse(str(doc_xml_path), parser)
    base_body = base_tree.getroot().find(f".//{{{W_NS}}}body")

    first_body = None
    for child in list(base_body):
        if child.tag == f"{{{W_NS}}}p" and text_of(child).strip().endswith("引言"):
            first_body = child
            break
    if first_body is None:
        raise RuntimeError("Cannot find template body start")
    insert_pos = list(base_body).index(first_body)
    for child in list(base_body)[insert_pos:]:
        if child.tag != f"{{{W_NS}}}sectPr":
            base_body.remove(child)

    with zipfile.ZipFile(content_docx) as cz:
        content_doc_xml = etree.fromstring(cz.read("word/document.xml"), parser)
        content_body = content_doc_xml.find(f".//{{{W_NS}}}body")
        content_rels_xml = etree.fromstring(cz.read("word/_rels/document.xml.rels"), parser)

        base_rels_path = temp / "word" / "_rels" / "document.xml.rels"
        base_rels = etree.parse(str(base_rels_path), parser)
        base_rels_root = base_rels.getroot()
        existing_ids = {rel.get("Id") for rel in base_rels_root}
        next_id = 1

        rid_map: dict[str, str] = {}
        media_dir = temp / "word" / "media"
        media_dir.mkdir(exist_ok=True)
        for rel in content_rels_xml:
            target = rel.get("Target", "")
            rel_type = rel.get("Type", "")
            old_id = rel.get("Id")
            if not target.startswith("media/"):
                continue
            while f"rIdMerged{next_id}" in existing_ids:
                next_id += 1
            new_id = f"rIdMerged{next_id}"
            existing_ids.add(new_id)
            next_id += 1
            src_name = "word/" + target
            ext = Path(target).suffix
            new_target = f"media/merged_{new_id}{ext}"
            (media_dir / Path(new_target).name).write_bytes(cz.read(src_name))
            new_rel = etree.Element(f"{{{PKG_REL_NS}}}Relationship")
            new_rel.set("Id", new_id)
            new_rel.set("Type", rel_type)
            new_rel.set("Target", new_target)
            base_rels_root.append(new_rel)
            rid_map[old_id] = new_id

        children = [child for child in list(content_body) if child.tag != f"{{{W_NS}}}sectPr"]
        for offset, child in enumerate(children):
            copied = copy.deepcopy(child)
            for elem in copied.xpath(".//*[@r:embed] | .//*[@r:id]", namespaces={"r": R_NS}):
                for attr in (f"{{{R_NS}}}embed", f"{{{R_NS}}}id"):
                    old = elem.get(attr)
                    if old in rid_map:
                        elem.set(attr, rid_map[old])
            base_body.insert(insert_pos + offset, copied)

        base_rels.write(str(base_rels_path), xml_declaration=True, encoding="UTF-8", standalone=True)

    base_tree.write(str(doc_xml_path), xml_declaration=True, encoding="UTF-8", standalone=True)
    content_types_path = temp / "[Content_Types].xml"
    ct_tree = etree.parse(str(content_types_path), parser)
    ct_root = ct_tree.getroot()
    ct_ns = ct_root.nsmap.get(None)
    existing_defaults = {elem.get("Extension") for elem in ct_root if elem.tag.endswith("Default")}
    for ext, ctype in {
        "png": "image/png",
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
    }.items():
        if ext not in existing_defaults:
            elem = etree.Element(f"{{{ct_ns}}}Default")
            elem.set("Extension", ext)
            elem.set("ContentType", ctype)
            ct_root.insert(0, elem)
    ct_tree.write(str(content_types_path), xml_declaration=True, encoding="UTF-8", standalone=True)
    with zipfile.ZipFile(out_docx, "w", zipfile.ZIP_DEFLATED) as zout:
        for path in temp.rglob("*"):
            if path.is_file():
                zout.write(path, path.relative_to(temp).as_posix())


def normalize_final_tables(path: Path) -> None:
    doc = Document(path)
    content_width = base.section_content_width_dxa(doc.sections[0])
    for idx, table in enumerate(doc.tables):
        cols = len(table.columns)
        if idx == 0 and cols == 2:
            widths = [1676, 5250]  # Preserve the template cover table's original proportions.
        elif cols == 2:
            widths = base.column_widths_from_weights([1.2, 3.8], content_width)
        elif cols == 5:
            widths = base.column_widths_from_weights([0.9, 1.5, 2.4, 2.4, 0.8], content_width)
        else:
            widths = base.column_widths_from_weights([1] * cols, content_width)
        base.apply_table_geometry(table, widths, table_width_dxa=sum(widths), indent_dxa=0)
    doc.save(path)


def main() -> None:
    WORK.mkdir(parents=True, exist_ok=True)
    front = WORK / "template-front.docx"
    shutil.copy(TEMPLATE, front)
    fill_template_front(front)

    source = build_content_source()
    content = WORK / "content.docx"
    subprocess.run(
        [
            "pandoc",
            str(source),
            "--from=latex",
            "--to=docx",
            f"--reference-doc={TEMPLATE}",
            f"--resource-path={base.WORK}:{WORK}:{THESIS}:{ROOT}",
            "-o",
            str(content),
        ],
        cwd=THESIS,
        check=True,
    )
    format_content_docx(content)
    merge_body_into_template(front, content, OUT)
    normalize_final_tables(OUT)
    base.patch_update_fields(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
