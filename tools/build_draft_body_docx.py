from __future__ import annotations

import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
THESIS = ROOT / "latex" / "thesis"
DRAFT = THESIS / "chapters" / "draft.tex"
BBL = THESIS / "build" / "main-body.bbl"
TEMPLATE = ROOT / "latex" / "hdu-template.docx"
WORK = Path("/private/tmp/hdu-draft-body-word")
OUT = ROOT / "draft正文手动迁移版.docx"

DOCS_SKILL = Path(
    "/Users/tk/.codex/plugins/cache/openai-primary-runtime/"
    "documents/26.430.10722/skills/documents/scripts"
)
sys.path.append(str(DOCS_SKILL))
from table_geometry import apply_table_geometry, column_widths_from_weights, section_content_width_dxa  # noqa: E402


def read(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def clean_latex_ref(text: str) -> str:
    replacements = {
        r"\newblock": " ",
        r"\allowbreak": "",
        r"\&": "&",
        r"\%": "%",
        r"\_": "_",
        r"---": "-",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"\\doi\{([^}]+)\}", r"https://doi.org/\1", text)
    text = re.sub(r"\\url\{([^}]+)\}", r"\1", text)
    text = re.sub(r"\\href\{([^}]+)\}\{([^}]+)\}", r"\2", text)
    text = re.sub(r"\\eprint\{([^}]+)\}\{([^}]+)\}", r"\2", text)
    text = re.sub(r"\\[a-zA-Z]+\*?(?:\[[^\]]*\])?", "", text)
    text = re.sub(r"[{}]", "", text)
    text = text.replace("~", " ")
    return re.sub(r"\s+", " ", text).strip()


def parse_bbl() -> tuple[dict[str, int], list[str]]:
    text = read(BBL)
    parts = re.split(r"\\bibitem(?:\[[^\]]*\])?\{([^}]+)\}", text)
    order: dict[str, int] = {}
    refs: list[str] = []
    for idx in range(1, len(parts), 2):
        key = parts[idx]
        raw = parts[idx + 1].split(r"\end{thebibliography}")[0]
        ref = clean_latex_ref(raw)
        if not ref:
            continue
        order[key] = len(refs) + 1
        refs.append(ref)
    return order, refs


def replace_citations(tex: str, order: dict[str, int]) -> str:
    def repl(match: re.Match[str]) -> str:
        nums = []
        for key in match.group(1).split(","):
            key = key.strip()
            if key in order:
                nums.append(str(order[key]))
        return "[" + ",".join(nums) + "]" if nums else ""

    return re.sub(r"\\cite\{([^}]+)\}", repl, tex)


def convert_pdf_to_png(src_pdf: Path, out_png: Path) -> bool:
    out_png.parent.mkdir(parents=True, exist_ok=True)
    magick = shutil.which("magick") or shutil.which("convert")
    try:
        if magick:
            subprocess.run(
                [
                    magick,
                    "-density",
                    "220",
                    str(src_pdf) + "[0]",
                    "-background",
                    "white",
                    "-alpha",
                    "remove",
                    "-alpha",
                    "off",
                    str(out_png),
                ],
                check=True,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
        else:
            subprocess.run(
                ["sips", "-s", "format", "png", str(src_pdf), "--out", str(out_png)],
                check=True,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
    except subprocess.CalledProcessError:
        return False
    return out_png.exists() and out_png.stat().st_size > 0


def prefer_word_safe_images(tex: str) -> str:
    def repl(match: re.Match[str]) -> str:
        options = match.group(1) or ""
        rel = Path(match.group(2))
        if rel.suffix.lower() != ".pdf":
            return match.group(0)

        pdf = THESIS / rel
        png_rel = rel.with_suffix(".png")
        thesis_png = THESIS / png_rel
        thesis_png_ok = thesis_png.exists() and thesis_png.stat().st_size > 0
        work_png = WORK / png_rel

        if not thesis_png_ok and pdf.exists():
            convert_pdf_to_png(pdf, work_png)

        if thesis_png_ok or (work_png.exists() and work_png.stat().st_size > 0):
            return f"\\includegraphics{options}" + "{" + str(png_rel) + "}"
        return match.group(0)

    return re.sub(r"\\includegraphics(\[[^\]]*\])?\{([^}]+)\}", repl, tex)


def build_source() -> Path:
    order, refs = parse_bbl()
    text = read(DRAFT)
    text = replace_citations(text, order)
    text = prefer_word_safe_images(text)

    ref_lines = ["\\chapter*{参考文献}"]
    for idx, ref in enumerate(refs, 1):
        ref_lines.append(f"\\noindent [{idx}] {ref}\n")

    source = text + "\n\n" + "\n".join(ref_lines)
    source = re.sub(r"\\label\{[^}]+\}", "", source)
    WORK.mkdir(parents=True, exist_ok=True)
    out = WORK / "draft-body.tex"
    out.write_text(source, encoding="utf-8")
    return out


def set_run_font(run, east_asia="宋体", ascii_font="Times New Roman", size=12, bold=None) -> None:
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("- ")
    set_run_font(run, size=10.5)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    run._r.append(instr)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)
    run = paragraph.add_run(" -")
    set_run_font(run, size=10.5)


def set_style_font(style, east_asia, ascii_font, size, bold=False, align=None):
    style.font.name = ascii_font
    style._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.paragraph_format.line_spacing = Pt(20)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    if align is not None:
        style.paragraph_format.alignment = align


def format_document(path: Path) -> None:
    doc = Document(path)
    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2)
        section.top_margin = Cm(3)
        section.bottom_margin = Cm(2)
        section.header_distance = Cm(1.5)
        section.footer_distance = Cm(1)
        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.text = "杭州电子科技大学继续教育学院本科毕业设计（论文）"
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in hp.runs:
            set_run_font(run, size=10.5)
        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.text = ""
        add_page_number(fp)

    set_style_font(doc.styles["Normal"], "宋体", "Times New Roman", 12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    doc.styles["Normal"].paragraph_format.first_line_indent = Pt(24)
    for name, size in [("Heading 1", 16), ("Heading 2", 15), ("Heading 3", 14)]:
        set_style_font(doc.styles[name], "黑体", "Times New Roman", size, True)
        doc.styles[name].paragraph_format.first_line_indent = Pt(0)
        doc.styles[name].paragraph_format.space_before = Pt(8)
        doc.styles[name].paragraph_format.space_after = Pt(8)
        doc.styles[name].paragraph_format.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER if name == "Heading 1" else WD_ALIGN_PARAGRAPH.LEFT
        )
    if "Body Text" in doc.styles:
        set_style_font(doc.styles["Body Text"], "宋体", "Times New Roman", 12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        doc.styles["Body Text"].paragraph_format.first_line_indent = Pt(24)
    if "First Paragraph" in doc.styles:
        set_style_font(doc.styles["First Paragraph"], "宋体", "Times New Roman", 12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        doc.styles["First Paragraph"].paragraph_format.first_line_indent = Pt(24)
    if "Caption" in doc.styles:
        set_style_font(doc.styles["Caption"], "宋体", "Times New Roman", 10.5, True, WD_ALIGN_PARAGRAPH.CENTER)
        doc.styles["Caption"].paragraph_format.first_line_indent = Pt(0)

    chapter = 0
    section = 0
    subsection = 0
    in_refs = False
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if p.style.name == "Heading 1":
            if text == "参考文献":
                p.text = text
                in_refs = True
            else:
                chapter += 1
                section = 0
                subsection = 0
                p.text = f"{chapter}  {text}"
            p.paragraph_format.first_line_indent = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if chapter > 1 or p.text == "参考文献":
                br = p.insert_paragraph_before("")
                br.paragraph_format.first_line_indent = Pt(0)
                br.add_run().add_break(WD_BREAK.PAGE)
        elif p.style.name == "Heading 2":
            section += 1
            subsection = 0
            p.text = f"{chapter}.{section}  {text}"
            p.paragraph_format.first_line_indent = Pt(0)
        elif p.style.name == "Heading 3":
            subsection += 1
            p.text = f"{chapter}.{section}.{subsection}  {text}"
            p.paragraph_format.first_line_indent = Pt(0)

        if re.match(r"^[图表]\s*\d", text):
            p.style = doc.styles["Caption"]
            p.paragraph_format.first_line_indent = Pt(0)

        for run in p.runs:
            if p.style.name.startswith("Heading"):
                size = 16 if p.style.name == "Heading 1" else 15 if p.style.name == "Heading 2" else 14
                set_run_font(run, "黑体", "Times New Roman", size, True)
            elif p.style.name == "Caption":
                set_run_font(run, "宋体", "Times New Roman", 10.5, True)
            elif in_refs:
                set_run_font(run, "宋体", "Times New Roman", 10.5)
            else:
                set_run_font(run, "宋体", "Times New Roman", 12)

    content_width = section_content_width_dxa(doc.sections[0])
    for table in doc.tables:
        cols = len(table.columns)
        if cols == 2:
            widths = column_widths_from_weights([1.2, 3.8], content_width)
        elif cols == 5:
            widths = column_widths_from_weights([0.9, 1.5, 2.4, 2.4, 0.8], content_width)
        else:
            widths = column_widths_from_weights([1] * cols, content_width)
        apply_table_geometry(table, widths, table_width_dxa=sum(widths), indent_dxa=0)

    doc.save(path)


def main() -> None:
    source = build_source()
    raw = WORK / "draft-body-raw.docx"
    subprocess.run(
        [
            "pandoc",
            str(source),
            "--from=latex",
            "--to=docx",
            f"--reference-doc={TEMPLATE}",
            f"--resource-path={WORK}:{THESIS}",
            "-o",
            str(raw),
        ],
        cwd=THESIS,
        check=True,
    )
    shutil.copy(raw, OUT)
    format_document(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
