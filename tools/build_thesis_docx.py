from __future__ import annotations

import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
THESIS = ROOT / "backup-202605011334" / "thesis"
TEMPLATE = ROOT / "backup-202605011334" / "hdu-template.docx"
WORK = Path("/private/tmp/hdu-thesis-word")
OUT = ROOT / "基于大模型的中小学生智能编程学习平台设计与实现-模板重制版.docx"
DOCS_SKILL = Path(
    "/Users/tk/.codex/plugins/cache/openai-primary-runtime/"
    "documents/26.430.10722/skills/documents/scripts"
)
sys.path.append(str(DOCS_SKILL))
from table_geometry import apply_table_geometry, column_widths_from_weights, section_content_width_dxa  # noqa: E402


def read(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def parse_metadata() -> dict[str, str]:
    text = read(THESIS / "metadata.tex")
    return dict(re.findall(r"\\newcommand\{\\([^}]+)\}\{([^}]*)\}", text))


def parse_bbl() -> tuple[dict[str, int], list[str]]:
    text = read(THESIS / "build" / "main.bbl")
    items = re.split(r"\\bibitem(?:\[[^\]]*\])?\{([^}]+)\}", text)
    order: dict[str, int] = {}
    refs: list[str] = []

    def clean_latex(s: str) -> str:
        replacements = {
            r"\newblock": " ",
            r"\allowbreak": "",
            r"\&": "&",
            r"\%": "%",
            r"\_": "_",
            r"\{": "{",
            r"\}": "}",
            r"---": "-",
        }
        for old, new in replacements.items():
            s = s.replace(old, new)
        s = re.sub(r"\\doi\{([^}]+)\}", r"https://doi.org/\1", s)
        s = re.sub(r"\\url\{([^}]+)\}", r"\1", s)
        s = re.sub(r"\\eprint\{([^}]+)\}\{([^}]+)\}", r"\2", s)
        s = re.sub(r"\\href\{([^}]+)\}\{([^}]+)\}", r"\2", s)
        s = re.sub(r"\\[a-zA-Z]+\*?(?:\[[^\]]*\])?", "", s)
        s = re.sub(r"[{}]", "", s)
        s = s.replace("~", " ")
        s = re.sub(r"\s+", " ", s)
        return s.strip()

    for i in range(1, len(items), 2):
        key = items[i]
        raw = items[i + 1]
        raw = raw.split(r"\end{thebibliography}")[0]
        ref = clean_latex(raw).strip()
        if not ref:
            continue
        order[key] = len(refs) + 1
        refs.append(ref)
    return order, refs


def replace_citations(tex: str, order: dict[str, int]) -> str:
    def repl(match: re.Match[str]) -> str:
        keys = [k.strip() for k in match.group(1).split(",") if k.strip()]
        nums = [str(order[k]) for k in keys if k in order]
        return "[" + ",".join(nums) + "]" if nums else ""

    return re.sub(r"\\cite\{([^}]+)\}", repl, tex)


def split_backmatter(text: str) -> tuple[str, str]:
    marker = r"\chapter*{致谢}"
    idx = text.find(marker)
    if idx < 0:
        return "", text
    return text[:idx], text[idx:]


def prefer_word_safe_images(tex: str) -> str:
    def repl(match: re.Match[str]) -> str:
        options = match.group(1) or ""
        rel = Path(match.group(2))
        if rel.suffix.lower() != ".pdf":
            return match.group(0)

        src_pdf = THESIS / rel
        png_rel = rel.with_suffix(".png")
        thesis_png = THESIS / png_rel
        work_png = WORK / png_rel
        thesis_png_ok = thesis_png.exists() and thesis_png.stat().st_size > 0
        if not thesis_png_ok and src_pdf.exists():
            work_png.parent.mkdir(parents=True, exist_ok=True)
            magick = shutil.which("magick") or shutil.which("convert")
            if magick:
                subprocess.run(
                    [
                        magick,
                        "-density",
                        "220",
                        str(src_pdf) + "[0]",
                        "-background",
                        "white",
                        "-alpha",
                        "remove",
                        "-alpha",
                        "off",
                        str(work_png),
                    ],
                    check=True,
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,
                )
            else:
                subprocess.run(
                    ["sips", "-s", "format", "png", str(src_pdf), "--out", str(work_png)],
                    check=True,
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,
                )
        if thesis_png_ok or (work_png.exists() and work_png.stat().st_size > 0):
            return f"\\includegraphics{options}" + "{" + str(png_rel) + "}"
        return match.group(0)

    return re.sub(r"\\includegraphics(\[[^\]]*\])?\{([^}]+)\}", repl, tex)


def build_merged_tex(order: dict[str, int], refs: list[str]) -> Path:
    meta = parse_metadata()
    cover = read(THESIS / "styles" / "cover.tex")
    abstract = read(THESIS / "chapters" / "abstract.tex")
    body = read(THESIS / "chapters" / "draft.tex")
    _, after_refs = split_backmatter(read(THESIS / "chapters" / "backmatter.tex"))

    replacements = {
        r"\ThesisTitleCn": meta["ThesisTitleCn"],
        r"\ThesisTitleEn": meta["ThesisTitleEn"],
        r"\ThesisCollege": meta["ThesisCollege"],
        r"\ThesisSite": meta["ThesisSite"],
        r"\ThesisMajor": meta["ThesisMajor"],
        r"\ThesisClass": meta["ThesisClass"],
        r"\ThesisStudentId": meta["ThesisStudentId"],
        r"\ThesisAuthor": meta["ThesisAuthor"],
        r"\ThesisAdvisor": meta["ThesisAdvisor"],
        r"\ThesisDate": meta["ThesisDate"],
    }

    text = "\n\n".join([cover, abstract, body])
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = replace_citations(text, order)

    ref_lines = ["\\chapter*{参考文献}"]
    for idx, ref in enumerate(refs, 1):
        ref_lines.append(f"\\noindent [{idx}] {ref}\n")

    after_refs = replace_citations(after_refs, order)
    for old, new in replacements.items():
        after_refs = after_refs.replace(old, new)

    merged = "\n\n".join([text, "\n".join(ref_lines), after_refs])
    merged = prefer_word_safe_images(merged)
    merged = re.sub(r"\\phantomsection|\\addcontentsline\{[^}]+\}\{[^}]+\}\{[^}]+\}|\\label\{[^}]+\}", "", merged)
    merged = merged.replace(r"\begin{CodeBlock}", r"\begin{verbatim}").replace(r"\end{CodeBlock}", r"\end{verbatim}")

    WORK.mkdir(parents=True, exist_ok=True)
    out = WORK / "merged.tex"
    out.write_text(merged, encoding="utf-8")
    return out


def set_cell_text(cell, text: str) -> None:
    cell.text = text
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = "楷体"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "楷体")
            run.font.size = Pt(15)


def set_run_font(run, east_asia="宋体", ascii_font="Times New Roman", size=12, bold=None) -> None:
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def ensure_style(doc: Document, name: str, base: str | None = None):
    try:
        style = doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    if base:
        style.base_style = doc.styles[base]
    return style


def set_paragraph_style(style, east_asia="宋体", ascii_font="Times New Roman", size=12, bold=False):
    font = style.font
    font.name = ascii_font
    font.size = Pt(size)
    font.bold = bold
    style._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for text in ("- ", " -"):
        if text.startswith("-"):
            run = paragraph.add_run(text)
            set_run_font(run, size=10.5)
            fld_begin = OxmlElement("w:fldChar")
            fld_begin.set(qn("w:fldCharType"), "begin")
            run._r.append(fld_begin)
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = "PAGE"
            run._r.append(instr)
            fld_end = OxmlElement("w:fldChar")
            fld_end.set(qn("w:fldCharType"), "end")
            run._r.append(fld_end)
        else:
            run = paragraph.add_run(text)
            set_run_font(run, size=10.5)


def add_toc(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-2" \h \z \u'
    run._r.append(instr)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run._r.append(separate)
    placeholder = OxmlElement("w:t")
    placeholder.text = "请在 Word 中右键更新目录"
    run._r.append(placeholder)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def format_document(path: Path) -> None:
    doc = Document(path)
    meta = parse_metadata()

    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2)
        section.top_margin = Cm(3)
        section.bottom_margin = Cm(2)
        section.header_distance = Cm(1.5)
        section.footer_distance = Cm(1)

    normal = doc.styles["Normal"]
    set_paragraph_style(normal, "宋体", "Times New Roman", 12)
    normal.paragraph_format.first_line_indent = Pt(24)
    normal.paragraph_format.line_spacing = Pt(20)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, bold, align in [
        ("Heading 1", 16, True, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 2", 15, True, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 3", 14, True, WD_ALIGN_PARAGRAPH.LEFT),
        ("Body Text", 12, False, WD_ALIGN_PARAGRAPH.JUSTIFY),
    ]:
        style = doc.styles[name]
        set_paragraph_style(style, "宋体" if name == "Body Text" else "黑体", "Times New Roman", size, bold)
        style.paragraph_format.line_spacing = Pt(20)
        style.paragraph_format.space_before = Pt(0 if name == "Body Text" else 8)
        style.paragraph_format.space_after = Pt(0 if name == "Body Text" else 8)
        style.paragraph_format.alignment = align
        if name == "Body Text":
            style.paragraph_format.first_line_indent = Pt(24)

    caption_style = ensure_style(doc, "Caption", "Normal")
    set_paragraph_style(caption_style, "宋体", "Times New Roman", 10.5, True)
    caption_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_style.paragraph_format.first_line_indent = Pt(0)

    ref_started = False
    chapter_no = 0
    section_no = 0
    subsection_no = 0
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue

        if text == "杭州电子科技大学":
            p.text = "杭州电子科技大学计算机学院"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                set_run_font(r, "黑体", "Times New Roman", 22, True)
            continue

        if text == "本科毕业设计（论文）":
            p.text = "本科毕业设计(论文）"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                set_run_font(r, "黑体", "Times New Roman", 22, True)
            continue

        if text.startswith("（2026"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                set_run_font(r, "宋体", "Times New Roman", 16)
            continue

        if text == meta["ThesisTitleCn"]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                set_run_font(r, "黑体", "Times New Roman", 16, True)
            continue

        if text in {"摘要", "摘    要"}:
            p.text = "摘    要"
            p.style = doc.styles["Body Text"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            for r in p.runs:
                set_run_font(r, "黑体", "Times New Roman", 16, True)
            continue

        if text == "ABSTRACT":
            p.style = doc.styles["Body Text"]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            for r in p.runs:
                set_run_font(r, "Times New Roman", "Times New Roman", 15, True)
            continue

        if text.startswith("关键词"):
            p.paragraph_format.first_line_indent = Pt(0)
            for r in p.runs:
                set_run_font(r, "宋体", "Times New Roman", 12)
            continue

        if text.lower().startswith("key words") or text.lower().startswith("key word"):
            p.text = text.replace("key words", "Key words").replace("key word", "Key words")
            p.paragraph_format.first_line_indent = Pt(0)
            for r in p.runs:
                set_run_font(r, "Times New Roman", "Times New Roman", 12)
            continue

        if text == "参考文献":
            ref_started = True

        if p.style.name == "Heading 1":
            if text in {"参考文献", "致谢", "附录"}:
                p.text = text
            else:
                chapter_no += 1
                section_no = 0
                p.text = f"{chapter_no}  {text}"
            p.paragraph_format.first_line_indent = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif p.style.name == "Heading 2":
            if text.startswith("附录 "):
                p.text = text
            else:
                section_no += 1
                subsection_no = 0
                p.text = f"{chapter_no}.{section_no}  {text}"
            p.paragraph_format.first_line_indent = Pt(0)
        elif p.style.name == "Heading 3":
            subsection_no += 1
            p.text = f"{chapter_no}.{section_no}.{subsection_no}  {text}"
            p.paragraph_format.first_line_indent = Pt(0)

        if text.startswith("图 ") or re.match(r"^图\s*\d", text) or text.startswith("表 "):
            p.style = caption_style
            p.paragraph_format.first_line_indent = Pt(0)

        for run in p.runs:
            if p.style.name.startswith("Heading"):
                set_run_font(run, "黑体", "Times New Roman", 16 if p.style.name == "Heading 1" else 15 if p.style.name == "Heading 2" else 14, True)
            elif p.style.name == "Caption":
                set_run_font(run, "宋体", "Times New Roman", 10.5, True)
            elif ref_started:
                set_run_font(run, "宋体", "Times New Roman", 10.5)
            else:
                set_run_font(run, "宋体", "Times New Roman", 12)

    page_break_before = {
        "诚 信 承 诺",
        "摘    要",
        "ABSTRACT",
        "目    录",
        "1  引言",
        "参考文献",
        "致谢",
        "附录",
    }
    for p in list(doc.paragraphs):
        if p.text.strip() in page_break_before:
            br = p.insert_paragraph_before("")
            br.paragraph_format.first_line_indent = Pt(0)
            br.add_run().add_break(WD_BREAK.PAGE)

    # Fill the cover metadata table converted from styles/cover.tex.
    if doc.tables:
        table = doc.tables[0]
        values = [
            ("站点名称：", meta["ThesisSite"]),
            ("专业：", meta["ThesisMajor"]),
            ("班级：", meta["ThesisClass"]),
            ("学号：", meta["ThesisStudentId"]),
            ("学生姓名：", meta["ThesisAuthor"]),
            ("指导教师：", meta["ThesisAdvisor"]),
            ("完成日期：", meta["ThesisDate"]),
        ]
        for row, (label, value) in zip(table.rows, values):
            set_cell_text(row.cells[0], label)
            set_cell_text(row.cells[1], value)

    content_width = section_content_width_dxa(doc.sections[0])
    for idx, table in enumerate(doc.tables):
        cols = len(table.columns)
        if cols == 2 and idx == 0:
            widths = column_widths_from_weights([1.2, 2.8], content_width)
        elif cols == 2:
            widths = column_widths_from_weights([1.2, 3.8], content_width)
        elif cols == 5:
            widths = column_widths_from_weights([0.9, 1.5, 2.4, 2.4, 0.8], content_width)
        else:
            widths = column_widths_from_weights([1] * cols, content_width)
        apply_table_geometry(table, widths, table_width_dxa=sum(widths), indent_dxa=0)

    # Insert a Word TOC field before the first main chapter.
    for idx, p in enumerate(doc.paragraphs):
        if p.text.strip().endswith("引言") and p.style.name == "Heading 1":
            toc_title = p.insert_paragraph_before("目    录")
            toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            toc_title.paragraph_format.first_line_indent = Pt(0)
            for r in toc_title.runs:
                set_run_font(r, "黑体", "Times New Roman", 16, True)
            toc = p.insert_paragraph_before("")
            add_toc(toc)
            break

    # Header/footer matching the school template wording.
    for section in doc.sections:
        header = section.header
        header.is_linked_to_previous = False
        hp = header.paragraphs[0]
        hp.text = "杭州电子科技大学继续教育学院本科毕业设计（论文）"
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in hp.runs:
            set_run_font(r, "宋体", "Times New Roman", 10.5)
        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0]
        fp.text = ""
        add_page_number(fp)

    doc.save(path)
    patch_toc_update(path)


def patch_toc_update(path: Path) -> None:
    tmp = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/settings.xml":
                xml = data.decode("utf-8")
                if "w:updateFields" not in xml:
                    xml = xml.replace("</w:settings>", '<w:updateFields w:val="true"/></w:settings>')
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


def main() -> None:
    order, refs = parse_bbl()
    merged = build_merged_tex(order, refs)
    raw_docx = WORK / "raw.docx"
    subprocess.run(
        [
            "pandoc",
            str(merged),
            "--from=latex",
            "--to=docx",
            f"--reference-doc={TEMPLATE}",
            f"--resource-path={WORK}:{THESIS}",
            "-o",
            str(raw_docx),
        ],
        cwd=THESIS,
        check=True,
    )
    shutil.copy(raw_docx, OUT)
    format_document(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
