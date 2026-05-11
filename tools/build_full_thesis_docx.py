from __future__ import annotations

import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
THESIS = ROOT / "latex" / "thesis"
TEMPLATE = ROOT / "latex" / "hdu-template.docx"
WORK = Path("/private/tmp/hdu-full-thesis-word")
OUT = ROOT / "论文完整手动迁移版.docx"

DOCS_SKILL = Path(
    "/Users/tk/.codex/plugins/cache/openai-primary-runtime/"
    "documents/26.430.10722/skills/documents/scripts"
)
sys.path.append(str(DOCS_SKILL))
from table_geometry import apply_table_geometry, column_widths_from_weights, section_content_width_dxa  # noqa: E402


def read(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def metadata() -> dict[str, str]:
    text = read(THESIS / "metadata.tex")
    return dict(re.findall(r"\\newcommand\{\\([^}]+)\}\{([^}]*)\}", text))


def clean_latex_ref(text: str) -> str:
    for old, new in {
        r"\newblock": " ",
        r"\allowbreak": "",
        r"\&": "&",
        r"\%": "%",
        r"\_": "_",
        r"---": "-",
    }.items():
        text = text.replace(old, new)
    text = re.sub(r"\\doi\{([^}]+)\}", r"https://doi.org/\1", text)
    text = re.sub(r"\\url\{([^}]+)\}", r"\1", text)
    text = re.sub(r"\\href\{([^}]+)\}\{([^}]+)\}", r"\2", text)
    text = re.sub(r"\\eprint\{([^}]+)\}\{([^}]+)\}", r"\2", text)
    text = re.sub(r"\\[a-zA-Z]+\*?(?:\[[^\]]*\])?", "", text)
    text = re.sub(r"[{}]", "", text).replace("~", " ")
    return re.sub(r"\s+", " ", text).strip()


def parse_bbl() -> tuple[dict[str, int], list[str]]:
    text = read(THESIS / "build" / "main-body.bbl")
    parts = re.split(r"\\bibitem(?:\[[^\]]*\])?\{([^}]+)\}", text)
    order: dict[str, int] = {}
    refs: list[str] = []
    for idx in range(1, len(parts), 2):
        key = parts[idx]
        raw = parts[idx + 1].split(r"\end{thebibliography}")[0]
        ref = clean_latex_ref(raw)
        if ref:
            order[key] = len(refs) + 1
            refs.append(ref)
    return order, refs


def replace_citations(tex: str, order: dict[str, int]) -> str:
    def repl(match: re.Match[str]) -> str:
        nums = [str(order[k.strip()]) for k in match.group(1).split(",") if k.strip() in order]
        return "[" + ",".join(nums) + "]" if nums else ""

    return re.sub(r"\\cite\{([^}]+)\}", repl, tex)


def convert_pdf_to_png(src_pdf: Path, out_png: Path) -> bool:
    out_png.parent.mkdir(parents=True, exist_ok=True)
    magick = shutil.which("magick") or shutil.which("convert")
    try:
        if magick:
            subprocess.run(
                [
                    magick,
                    "-density",
                    "220",
                    str(src_pdf) + "[0]",
                    "-background",
                    "white",
                    "-alpha",
                    "remove",
                    "-alpha",
                    "off",
                    str(out_png),
                ],
                check=True,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
        else:
            subprocess.run(
                ["sips", "-s", "format", "png", str(src_pdf), "--out", str(out_png)],
                check=True,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
    except subprocess.CalledProcessError:
        return False
    return out_png.exists() and out_png.stat().st_size > 0


def prefer_word_safe_images(tex: str) -> str:
    def repl(match: re.Match[str]) -> str:
        opts = match.group(1) or ""
        rel = Path(match.group(2))
        if rel.suffix.lower() != ".pdf":
            return match.group(0)
        pdf = (THESIS / rel).resolve()
        png_rel = rel.with_suffix(".png")
        thesis_png = THESIS / png_rel
        thesis_png_ok = thesis_png.exists() and thesis_png.stat().st_size > 0
        work_png = WORK / png_rel
        if not thesis_png_ok and pdf.exists():
            convert_pdf_to_png(pdf, work_png)
        if thesis_png_ok or (work_png.exists() and work_png.stat().st_size > 0):
            return f"\\includegraphics{opts}" + "{" + str(png_rel) + "}"
        return match.group(0)

    return re.sub(r"\\includegraphics(\[[^\]]*\])?\{([^}]+)\}", repl, tex)


def strip_latex_inline(text: str) -> str:
    text = re.sub(r"\\(?:zihao|fontspec|CJKfontspec)(?:\[[^\]]*\])?\{[^}]*\}", "", text)
    text = re.sub(r"\\[A-Za-z]+\*?(?:\[[^\]]*\])?\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\[A-Za-z]+\*?(?:\[[^\]]*\])?", "", text)
    text = re.sub(r"[{}]", "", text)
    return re.sub(r"\s+", " ", text.replace(r"\&", "&")).strip()


def abstract_parts() -> tuple[list[str], str, list[str], str]:
    lines = read(THESIS / "chapters" / "abstract.tex").splitlines()

    cn_raw: list[str] = []
    en_raw: list[str] = []
    cn_kw = ""
    en_kw = ""
    state: str | None = None
    for line in lines:
        stripped = line.strip()
        if stripped.startswith(r"{\song"):
            state = "cn"
            continue
        if stripped.startswith(r"{\fontspec{Times New Roman}"):
            state = "en"
            continue
        if state and stripped.startswith(r"\vspace{20pt}"):
            state = f"{state}_kw"
            continue
        if "关键词" in stripped:
            cn_kw = strip_latex_inline(stripped)
            cn_kw = re.sub(r"^.*?关键词[：:]", "", cn_kw).strip()
            state = None
            continue
        if "Key Words" in stripped or "key words" in stripped:
            en_kw = strip_latex_inline(stripped)
            en_kw = re.sub(r"^.*?Key Words[：:]", "", en_kw, flags=re.I).strip()
            en_kw = re.sub(r"^(?:Key\s+)?Words[：:]\s*", "", en_kw, flags=re.I).strip()
            state = None
            continue
        if state == "cn":
            cn_raw.append(line)
        elif state == "en":
            en_raw.append(line)

    cn_body = [strip_latex_inline(p) for p in re.split(r"\n\s*\n", "\n".join(cn_raw)) if strip_latex_inline(p)]
    en_body = [strip_latex_inline(p) for p in re.split(r"\n\s*\n", "\n".join(en_raw)) if strip_latex_inline(p)]
    return cn_body, cn_kw, en_body, en_kw


def backmatter_after_refs(order: dict[str, int]) -> str:
    text = read(THESIS / "chapters" / "backmatter.tex")
    idx = text.find(r"\chapter*{致谢}")
    if idx == -1:
        return ""
    text = text[idx:]
    text = replace_citations(text, order)
    text = prefer_word_safe_images(text)
    text = re.sub(r"\\addcontentsline\{[^}]+\}\{[^}]+\}\{[^}]+\}|\\phantomsection|\\label\{[^}]+\}", "", text)
    return text


def build_source() -> Path:
    meta = metadata()
    order, refs = parse_bbl()
    cn_body, cn_kw, en_body, en_kw = abstract_parts()
    draft = replace_citations(read(THESIS / "chapters" / "draft.tex"), order)
    draft = prefer_word_safe_images(draft)
    draft = re.sub(r"\\label\{[^}]+\}", "", draft)

    cover = rf"""
\begin{{center}}
\textbf{{杭州电子科技大学计算机学院}}

\textbf{{本科毕业设计(论文）}}

（2026届）
\end{{center}}

\begin{{tabular}}{{ll}}
题\quad 目 & {meta['ThesisTitleCn']} \\
站点名称 & {meta['ThesisSite']} \\
专\quad 业 & {meta['ThesisMajor']} \\
班\quad 级 & {meta['ThesisClass']} \\
学\quad 号 & {meta['ThesisStudentId']} \\
学生姓名 & {meta['ThesisAuthor']} \\
指导教师 & {meta['ThesisAdvisor']} \\
完成日期 & {meta['ThesisDate']} \\
\end{{tabular}}

\newpage
\begin{{center}}\textbf{{诚 信 承 诺}}\end{{center}}

我谨在此承诺：本人所写的毕业论文《{meta['ThesisTitleCn']}》均系本人独立完成，没有抄袭行为，凡涉及其他作者的观点和材料，均作了注释，若有不实，后果由本人承担。

\begin{{flushright}}
承诺人（签名）：\underline{{\hspace{{5cm}}}}\\
年\quad 月\quad 日
\end{{flushright}}
"""
    abstract = ["\\newpage", "\\begin{center}\\textbf{摘    要}\\end{center}", *cn_body, f"\\noindent \\textbf{{关键词：}}{cn_kw}"]
    abstract += ["\\newpage", "\\begin{center}\\textbf{ABSTRACT}\\end{center}", *en_body, f"\\noindent \\textbf{{Key Words：}}{en_kw}"]
    refs_tex = ["\\chapter*{参考文献}"] + [f"\\noindent [{idx}] {ref}\n" for idx, ref in enumerate(refs, 1)]
    source = "\n\n".join([cover, "\n\n".join(abstract), draft, "\n".join(refs_tex), backmatter_after_refs(order)])

    WORK.mkdir(parents=True, exist_ok=True)
    out = WORK / "full-thesis.tex"
    out.write_text(source, encoding="utf-8")
    return out


def set_run_font(run, east_asia="宋体", ascii_font="Times New Roman", size=12, bold=None) -> None:
    run.font.name = ascii_font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def style_font(style, east_asia, ascii_font, size, bold=False, align=None):
    style.font.name = ascii_font
    style._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.paragraph_format.line_spacing = Pt(20)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = Pt(0)
    if align is not None:
        style.paragraph_format.alignment = align


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("- ")
    set_run_font(run, size=10.5)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    run._r.append(instr)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)
    run = paragraph.add_run(" -")
    set_run_font(run, size=10.5)


def add_toc(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    run._r.append(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-2" \h \z \u'
    run._r.append(instr)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.append(sep)
    t = OxmlElement("w:t")
    t.text = "请在 Word 中右键更新目录"
    run._r.append(t)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def patch_update_fields(path: Path) -> None:
    tmp = path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/settings.xml":
                xml = data.decode("utf-8")
                if "w:updateFields" not in xml:
                    xml = xml.replace("</w:settings>", '<w:updateFields w:val="true"/></w:settings>')
                data = xml.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


def format_docx(path: Path) -> None:
    doc = Document(path)
    for sec in doc.sections:
        sec.page_width = Cm(21)
        sec.page_height = Cm(29.7)
        sec.left_margin = Cm(3)
        sec.right_margin = Cm(2)
        sec.top_margin = Cm(3)
        sec.bottom_margin = Cm(2)
        sec.header_distance = Cm(1.5)
        sec.footer_distance = Cm(1)
        sec.header.is_linked_to_previous = False
        hp = sec.header.paragraphs[0]
        hp.text = "杭州电子科技大学继续教育学院本科毕业设计（论文）"
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in hp.runs:
            set_run_font(run, size=10.5)
        sec.footer.is_linked_to_previous = False
        fp = sec.footer.paragraphs[0]
        fp.text = ""
        add_page_number(fp)

    style_font(doc.styles["Normal"], "宋体", "Times New Roman", 12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    doc.styles["Normal"].paragraph_format.first_line_indent = Pt(24)
    if "Body Text" in doc.styles:
        style_font(doc.styles["Body Text"], "宋体", "Times New Roman", 12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        doc.styles["Body Text"].paragraph_format.first_line_indent = Pt(24)
    if "First Paragraph" in doc.styles:
        style_font(doc.styles["First Paragraph"], "宋体", "Times New Roman", 12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        doc.styles["First Paragraph"].paragraph_format.first_line_indent = Pt(24)
    for name, size in [("Heading 1", 16), ("Heading 2", 15), ("Heading 3", 14)]:
        style_font(doc.styles[name], "黑体", "Times New Roman", size, True)
        doc.styles[name].paragraph_format.first_line_indent = Pt(0)
        doc.styles[name].paragraph_format.space_before = Pt(8)
        doc.styles[name].paragraph_format.space_after = Pt(8)
        doc.styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if name == "Heading 1" else WD_ALIGN_PARAGRAPH.LEFT
    if "Caption" in doc.styles:
        style_font(doc.styles["Caption"], "宋体", "Times New Roman", 10.5, True, WD_ALIGN_PARAGRAPH.CENTER)
        doc.styles["Caption"].paragraph_format.first_line_indent = Pt(0)

    chapter = section = subsection = 0
    specials = {"参考文献", "致谢", "附录"}
    in_refs = False
    first_body_heading = None
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if text in {"摘 要", "摘    要", "ABSTRACT", "诚 信 承 诺"}:
            if text in {"摘 要", "摘    要"}:
                p.text = "摘    要"
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            for run in p.runs:
                set_run_font(run, "黑体" if text != "ABSTRACT" else "Times New Roman", "Times New Roman", 16, True)
        if text == "杭州电子科技大学计算机学院":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            for run in p.runs:
                set_run_font(run, "黑体", "Times New Roman", 22, True)
        if text == "本科毕业设计(论文）":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            for run in p.runs:
                set_run_font(run, "黑体", "Times New Roman", 22, True)
        if text == "（2026届）":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
        if p.style.name == "Heading 1":
            if text in specials:
                p.text = text
                in_refs = text == "参考文献"
            else:
                chapter += 1
                section = subsection = 0
                p.text = f"{chapter}  {text}"
                if first_body_heading is None:
                    first_body_heading = p
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            if chapter > 1 or text in specials:
                br = p.insert_paragraph_before("")
                br.paragraph_format.first_line_indent = Pt(0)
                br.add_run().add_break(WD_BREAK.PAGE)
        elif p.style.name == "Heading 2":
            if text.startswith("附录 "):
                p.text = text
            else:
                section += 1
                subsection = 0
                p.text = f"{chapter}.{section}  {text}"
            p.paragraph_format.first_line_indent = Pt(0)
        elif p.style.name == "Heading 3":
            subsection += 1
            p.text = f"{chapter}.{section}.{subsection}  {text}"
            p.paragraph_format.first_line_indent = Pt(0)
        if re.match(r"^[图表]\s*\d", text):
            p.style = doc.styles["Caption"]
            p.paragraph_format.first_line_indent = Pt(0)
        for run in p.runs:
            if p.style.name.startswith("Heading"):
                size = 16 if p.style.name == "Heading 1" else 15 if p.style.name == "Heading 2" else 14
                set_run_font(run, "黑体", "Times New Roman", size, True)
            elif p.style.name == "Caption":
                set_run_font(run, "宋体", "Times New Roman", 10.5, True)
            elif in_refs:
                set_run_font(run, "宋体", "Times New Roman", 10.5)
            else:
                set_run_font(run, "宋体", "Times New Roman", 12)

    if first_body_heading is not None:
        title = first_body_heading.insert_paragraph_before("目    录")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title.paragraph_format.first_line_indent = Pt(0)
        for r in title.runs:
            set_run_font(r, "黑体", "Times New Roman", 16, True)
        toc = first_body_heading.insert_paragraph_before("")
        toc.paragraph_format.first_line_indent = Pt(0)
        add_toc(toc)
        br = first_body_heading.insert_paragraph_before("")
        br.paragraph_format.first_line_indent = Pt(0)
        br.add_run().add_break(WD_BREAK.PAGE)

    if doc.tables:
        table = doc.tables[0]
        meta = metadata()
        values = [
            ("题    目", meta["ThesisTitleCn"]),
            ("站点名称", meta["ThesisSite"]),
            ("专    业", meta["ThesisMajor"]),
            ("班    级", meta["ThesisClass"]),
            ("学    号", meta["ThesisStudentId"]),
            ("学生姓名", meta["ThesisAuthor"]),
            ("指导教师", meta["ThesisAdvisor"]),
            ("完成日期", meta["ThesisDate"]),
        ]
        for row, (label, value) in zip(table.rows, values):
            row.cells[0].text = label
            row.cells[1].text = value
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.first_line_indent = Pt(0)
                    for r in p.runs:
                        set_run_font(r, "楷体", "Times New Roman", 15)

    content_width = section_content_width_dxa(doc.sections[0])
    for idx, table in enumerate(doc.tables):
        cols = len(table.columns)
        if idx == 0 and cols == 2:
            widths = column_widths_from_weights([1.2, 3.8], content_width)
        elif cols == 2:
            widths = column_widths_from_weights([1.2, 3.8], content_width)
        elif cols == 5:
            widths = column_widths_from_weights([0.9, 1.5, 2.4, 2.4, 0.8], content_width)
        else:
            widths = column_widths_from_weights([1] * cols, content_width)
        apply_table_geometry(table, widths, table_width_dxa=sum(widths), indent_dxa=0)

    doc.save(path)
    patch_update_fields(path)


def main() -> None:
    source = build_source()
    raw = WORK / "full-thesis-raw.docx"
    subprocess.run(
        [
            "pandoc",
            str(source),
            "--from=latex",
            "--to=docx",
            f"--reference-doc={TEMPLATE}",
            f"--resource-path={WORK}:{THESIS}:{ROOT}",
            "-o",
            str(raw),
        ],
        cwd=THESIS,
        check=True,
    )
    shutil.copy(raw, OUT)
    format_docx(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
